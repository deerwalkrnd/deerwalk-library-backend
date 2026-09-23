import csv
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO, StringIO
from typing import Any, Dict, List, Optional
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

from app.core.exc.error_code import ErrorCode
from app.core.exc.library_exception import LibraryException

# How far down a sheet to look for the header row. The library's register has
# sheets whose first row is blank, with the real header on row 2.
HEADER_SEARCH_ROWS = 10

# What openpyxl hands back for formula errors, and for cells formatted as a
# date whose number is too large to be one.
EXCEL_ERROR_VALUES = {
    "#N/A",
    "#NAME?",
    "#NULL!",
    "#NUM!",
    "#REF!",
    "#VALUE!",
    "#DIV/0!",
}


@dataclass
class SheetRow:
    number: int  # 1-indexed row number as shown in Excel
    values: Dict[str, Any]


@dataclass
class Sheet:
    name: Optional[str]  # None for CSV files
    headers: List[str]  # normalised, see normalize_header
    rows: List[SheetRow] = field(default_factory=list)


def normalize_header(value: Any) -> str:
    """'  Accession Number ' -> 'accession number'."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip().lower()


def cell_to_text(value: Any) -> Optional[str]:
    """
    Render a spreadsheet cell as the text a librarian typed.

    Excel stores ISBNs and accession numbers as floats (9781409550525.0) and
    occasionally auto-formats a plain number as a date; both are turned back
    into the digits that were entered.
    """
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    if isinstance(value, (datetime, date)):
        serial = to_excel(value)
        return str(int(serial)) if float(serial).is_integer() else str(serial)
    text = re.sub(r"\s+", " ", str(value)).strip()
    if text in EXCEL_ERROR_VALUES:
        return None
    return text or None


def read_sheets(filename: str, content: bytes) -> List[Sheet]:
    name = filename.lower()
    if name.endswith(".csv"):
        return [_read_csv(content)]
    if name.endswith(".xlsx"):
        return _read_xlsx(content)
    if name.endswith(".xls"):
        raise LibraryException(
            status_code=400,
            code=ErrorCode.INVALID_FIELDS,
            msg="Old .xls files are not supported. Open the file in Excel and save it as .xlsx.",
        )
    raise LibraryException(
        status_code=400,
        code=ErrorCode.INVALID_FIELDS,
        msg="Only CSV and Excel (.xlsx) files are allowed!",
    )


def _read_csv(content: bytes) -> Sheet:
    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        decoded = content.decode("latin-1")

    return _build_sheet(name=None, raw_rows=list(csv.reader(StringIO(decoded))))


def _read_xlsx(content: bytes) -> List[Sheet]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, KeyError, OSError, ValueError) as e:
        raise LibraryException(
            status_code=400,
            code=ErrorCode.INVALID_FIELDS,
            msg=f"Could not read the Excel file: {e}",
        )

    try:
        return [
            _build_sheet(
                name=worksheet.title.strip(),
                raw_rows=[list(row) for row in worksheet.iter_rows(values_only=True)],
            )
            for worksheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def _build_sheet(name: Optional[str], raw_rows: List[List[Any]]) -> Sheet:
    header_index = next(
        (
            i
            for i, row in enumerate(raw_rows[:HEADER_SEARCH_ROWS])
            if "title" in {normalize_header(cell) for cell in row}
        ),
        None,
    )
    if header_index is None:
        # Not a book list (e.g. the register's summary "Report" sheet).
        return Sheet(name=name, headers=[])

    headers = [normalize_header(cell) for cell in raw_rows[header_index]]
    sheet = Sheet(name=name, headers=[h for h in headers if h])

    for offset, raw in enumerate(raw_rows[header_index + 1 :]):
        values: Dict[str, Any] = {}
        for header, cell in zip(headers, raw):
            # First occurrence wins if a header is repeated.
            if header and header not in values:
                values[header] = cell
        if any(cell_to_text(v) for v in values.values()):
            sheet.rows.append(
                SheetRow(number=header_index + offset + 2, values=values)
            )

    return sheet
