import csv
from io import BytesIO, StringIO
from typing import Iterable, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.modules.books.utils.book_import.register_format import ACADEMIC_SUBJECTS

# The library's accession register layout (see register_format.py), so the
# template is the sheet librarians already keep. One row per physical copy.
REGISTER_TEMPLATE_HEADERS = [
    "SN.",
    "Date",
    "Acc no",
    "AUTHOR1",
    "TITLE",
    "PLACE AND PUBLISHER",
    "Year",
    "Pages",
    "PRICE",
    "ISBN NO.",
    "CONDITION",
    "AUTHOR2",
    "TITLE1",
    "CATEGORY",
]
COLUMN_WIDTHS = [6, 12, 8, 24, 36, 28, 8, 8, 10, 18, 12, 20, 24, 14]

# The sheet name is part of every copy's ID (see register_format.sheet_prefix).
TEMPLATE_SHEET_NAME = "GENERAL"

# (column, required, what to enter, example)
COLUMN_GUIDE = [
    ("SN.", "No", "Serial number. Not imported.", "1"),
    ("Date", "No", "Date the copy was registered. Not imported.", "2079-04-23"),
    (
        "Acc no",
        "Yes",
        "Accession number of this copy. Must be unique within the sheet. "
        "The copy is saved as <SHEET NAME>-<Acc no>, e.g. GENERAL-101.",
        "101",
    ),
    ("AUTHOR1", "No", "Main author.", "Julia Donaldson"),
    ("TITLE", "Yes", "Book title.", "What the Ladybird Heard"),
    ("PLACE AND PUBLISHER", "No", "Publisher.", "London: Macmillan"),
    ("Year", "No", "Not imported.", "2017"),
    ("Pages", "No", "Not imported.", "32"),
    ("PRICE", "No", "Not imported.", "Rs 480"),
    (
        "ISBN NO.",
        "No",
        "ISBN-10 or ISBN-13; dashes and spaces are fine. Rows with the same "
        "ISBN (or, without one, the same title and author) become copies of "
        "one book.",
        "9781447208648",
    ),
    ("CONDITION", "No", "Condition of this copy.", "Good"),
    ("AUTHOR2", "No", "Second author, if any.", ""),
    ("TITLE1", "No", "Subtitle. Used as the title only when TITLE is empty.", ""),
    (
        "CATEGORY",
        "No",
        "Subject. These make the book ACADEMIC: "
        + ", ".join(sorted(ACADEMIC_SUBJECTS))
        + ". Anything else (or empty) is NON_ACADEMIC. Dictionaries, "
        "encyclopedias, atlases and thesauruses become REFERENCE from their title.",
        "Fiction",
    ),
]

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1E78B9")


def build_xlsx_template(genre_names: Iterable[str]) -> bytes:
    workbook = Workbook()

    books = workbook.active
    books.title = TEMPLATE_SHEET_NAME
    books.append(REGISTER_TEMPLATE_HEADERS)
    _style_header(books, widths=COLUMN_WIDTHS)
    books.freeze_panes = "A2"
    # Keep accession numbers and ISBNs as typed (no 9.78E+12, no dropped zeros).
    for header in ("Acc no", "ISBN NO."):
        column = _column_letter(REGISTER_TEMPLATE_HEADERS.index(header))
        books.column_dimensions[column].number_format = "@"

    # One column name per row, so the importer never mistakes this sheet
    # (whose rows contain "TITLE" but never alongside "AUTHOR1"/"Acc no") for data.
    guide = workbook.create_sheet("Instructions")
    guide.append(["Column", "Required", "What to enter", "Example"])
    for row in COLUMN_GUIDE:
        guide.append(list(row))
    guide.append([])
    for note in (
        "One row per physical copy: three copies of a book are three rows with "
        "three accession numbers.",
        f"Rename the {TEMPLATE_SHEET_NAME} sheet, or add more sheets with the "
        "same columns, to keep separate accession series (e.g. FICTION ENG, "
        "NEPALI MSHS). Each sheet's name becomes the start of its copies' IDs.",
        "Copies whose ID is already in the library are skipped, so the same "
        "file can be uploaded again after adding rows.",
        "Genres and grade are not in the register; set them on the book page "
        "after importing. Existing genres are listed on the Genres sheet.",
    ):
        guide.append(["Note", "", note])
    _style_header(guide, widths=[22, 10, 80, 30])
    for row in guide.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    genres = workbook.create_sheet("Genres")
    genres.append(["Available genres"])
    for name in sorted(genre_names, key=str.casefold):
        genres.append([name])
    _style_header(genres, widths=[32])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_csv_template() -> bytes:
    buffer = StringIO()
    csv.writer(buffer).writerow(REGISTER_TEMPLATE_HEADERS)
    # BOM so Excel opens the file as UTF-8 (Nepali titles).
    return buffer.getvalue().encode("utf-8-sig")


def _column_letter(index: int) -> str:
    return chr(ord("A") + index)


def _style_header(sheet, widths: List[int]) -> None:
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for index, width in enumerate(widths):
        sheet.column_dimensions[_column_letter(index)].width = width
